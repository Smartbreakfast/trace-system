# -*- coding: utf-8 -*-
"""Đổi một tệp Markdown trong tai-lieu/ thành .docx đúng khổ của báo cáo.

    python tools/md_sang_docx.py tai-lieu/SO_SANH_CONG_NGHE.md

Ra tệp .docx cùng tên, cùng thư mục. Khổ theo quy định hồ sơ dự thi khoa học kỹ
thuật: A4, lề trái 3cm, phải 2cm, trên 2cm, dưới 2cm, Times New Roman 14, cách
dòng đơn. Dán từ tệp này sang báo cáo là không phải chỉnh lại định dạng.

Chỉ hiểu phần Markdown mà tài liệu trong tai-lieu/ dùng: tiêu đề `#`, đoạn văn,
danh sách `-` và `1.`, bảng `|`, khối mã ``` và chữ in đậm, chữ mã trong đoạn.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = 'Times New Roman'
CODE_FONT = 'Consolas'
BODY_SIZE = Pt(14)
TABLE_SIZE = Pt(11)
CODE_SIZE = Pt(10.5)


def setup(document):
    style = document.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    # Chữ tiếng Việt lấy font từ thuộc tính eastAsia khi Word đoán sai bộ chữ.
    style.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    paragraph = style.paragraph_format
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.0

    for section in document.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)


def inline(paragraph, text, size=BODY_SIZE, bold=False):
    """Viết một đoạn chữ, hiểu **in đậm** và `chữ mã`."""
    for piece in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(size.pt - 2)
        else:
            run = paragraph.add_run(piece)
            run.bold = bold
        run.font.size = size


def heading(document, text, level):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level < 3 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(re.sub(r'\*\*', '', text))
    run.bold = True
    run.italic = level >= 3
    run.font.name = BODY_FONT
    run.font.size = Pt({1: 16, 2: 14, 3: 14}.get(level, 14))


def table(document, rows):
    columns = max(len(row) for row in rows)
    created = document.add_table(rows=0, cols=columns)
    created.style = 'Table Grid'
    created.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, row in enumerate(rows):
        cells = created.add_row().cells
        for column in range(columns):
            cell = cells[column]
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            inline(
                cell.paragraphs[0],
                row[column] if column < len(row) else '',
                size=TABLE_SIZE,
                bold=index == 0,
            )
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def code_block(document, lines):
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line)
        run.font.name = CODE_FONT
        run.font.size = CODE_SIZE


def split_row(line):
    return [cell.strip() for cell in line.strip().strip('|').split('|')]


def convert(source, target):
    document = Document()
    setup(document)

    with open(source, encoding='utf-8') as handle:
        lines = handle.read().split('\n')

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped or stripped == '---':
            index += 1
            continue

        if stripped.startswith('```'):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                block.append(lines[index])
                index += 1
            code_block(document, block)
            index += 1
            continue

        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            heading(document, stripped[level:].strip(), level)
            index += 1
            continue

        if stripped.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = lines[index].strip()
                # Dòng gạch ngang tách phần đầu bảng, không phải dữ liệu.
                if not re.fullmatch(r'\|[\s:|-]+\|', row):
                    rows.append(split_row(row))
                index += 1
            if rows:
                table(document, rows)
            continue

        bullet = re.match(r'^(\s*)[-*]\s+(.*)', line)
        number = re.match(r'^(\s*)\d+\.\s+(.*)', line)
        if bullet or number:
            match = bullet or number
            depth = len(match.group(1)) // 2
            text = match.group(2)
            # Dòng tiếp của cùng một gạch đầu dòng được thụt vào, gộp lại.
            while (
                index + 1 < len(lines)
                and lines[index + 1].startswith('  ')
                and lines[index + 1].strip()
                and not re.match(r'^\s*([-*]|\d+\.)\s', lines[index + 1])
                and not lines[index + 1].strip().startswith('|')
            ):
                index += 1
                text += ' ' + lines[index].strip()
            style = 'List Bullet' if bullet else 'List Number'
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.left_indent = Cm(0.75 + depth * 0.75)
            paragraph.paragraph_format.space_after = Pt(3)
            inline(paragraph, text)
            index += 1
            continue

        text = stripped
        while (
            index + 1 < len(lines)
            and lines[index + 1].strip()
            and not re.match(r'^\s*([-*#|]|\d+\.|```)', lines[index + 1])
        ):
            index += 1
            text += ' ' + lines[index].strip()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        inline(paragraph, text)
        index += 1

    document.save(target)
    return target


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    source = sys.argv[1]
    target = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(source)[0] + '.docx'
    convert(source, target)
    size = os.path.getsize(target)
    print('đã ghi %s (%.1f KB)' % (target, size / 1024))


if __name__ == '__main__':
    main()

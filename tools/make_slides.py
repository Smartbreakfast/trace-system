# -*- coding: utf-8 -*-
"""Sinh bộ slide trình bày Tú Lệ Trace.

Chạy:  python tools/make_slides.py
Kết quả: tai-lieu/Tu-Le-Trace-Slides.pptx

Nội dung slide bám theo tai-lieu/KIEN_TRUC.md, tai-lieu/HUONG_DAN_SU_DUNG.md
và QUY_TRINH.md. Sửa nội dung thì sửa ở đây rồi chạy lại, đừng sửa tay file
pptx vì lần chạy sau sẽ ghi đè.

Ảnh minh hoạ phần hướng dẫn sử dụng lấy trong tai-lieu/anh/ theo đúng tên file
ghi ở tai-lieu/anh/README.md. File nào chưa có thì slide hiện khung chờ ảnh;
chụp xong chạy lại script là ảnh vào đúng chỗ.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'tai-lieu', 'Tu-Le-Trace-Slides.pptx')
SHOT = os.path.join(ROOT, 'tai-lieu', 'anh')
ASSET = os.path.join(ROOT, 'assets')

INK = RGBColor(0x17, 0x35, 0x2A)
INK_SOFT = RGBColor(0x2C, 0x4C, 0x3E)
ORANGE = RGBColor(0xE2, 0x70, 0x3A)
BG = RGBColor(0xF6, 0xF6, 0xF1)
BG2 = RGBColor(0xEE, 0xF0, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xDA, 0xCC)
HAIR = RGBColor(0xE2, 0xE5, 0xDA)
MUTED = RGBColor(0x5C, 0x6B, 0x61)
PALE = RGBColor(0xC4, 0xD6, 0xCA)
GREEN_SOFT = RGBColor(0xEF, 0xF7, 0xEC)
GREEN_EDGE = RGBColor(0xCB, 0xE3, 0xC0)

FONT = 'Segoe UI'
W, H = Inches(13.333), Inches(7.5)
ML = Inches(0.85)
CW = Inches(11.63)
TOP = Inches(1.9)
HALF = int((CW - Inches(0.4)) / 2)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
_n = [0]


def _txt(shape, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space=6, line=1.25):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    if isinstance(runs, str):
        runs = [runs]
    for i, item in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line
        text_, opts = item if isinstance(item, tuple) else (item, {})
        r = p.add_run()
        r.text = text_
        f = r.font
        f.name = FONT
        f.size = Pt(opts.get('size', size))
        f.bold = opts.get('bold', bold)
        f.color.rgb = opts.get('color', color)
    return tf


def box(slide, left, top, width, height, fill=None, edge=None, radius=False,
        dash=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        int(left), int(top), int(width), int(height))
    if radius:
        shape.adjustments[0] = 0.06
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if edge is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = edge
        shape.line.width = Pt(1.0)
        if dash:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def text(slide, left, top, width, height, runs, **kw):
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    _txt(tb, runs, **kw)
    return tb


def label(s, left, top, width, txt, color=ORANGE):
    """Nhãn nhỏ in hoa cho tiêu đề khối bên trong slide."""
    return text(s, left, top, width, Inches(0.26), txt.upper(),
                size=11, color=color, bold=True, space=0)


def slide(title_text, section=None):
    """Tiêu đề và gạch cam. Tên phần đặt ở chân trang, không lặp trên đầu."""
    _n[0] += 1
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, W, H, fill=BG)
    text(s, ML, Inches(0.72), CW, Inches(0.62), title_text,
         size=27, color=INK, bold=True, space=0)
    box(s, ML, Inches(1.46), Inches(1.15), Pt(2.6), fill=ORANGE)
    text(s, W - Inches(1.6), H - Inches(0.62), Inches(0.9), Inches(0.3),
         str(_n[0]), size=10, color=MUTED, align=PP_ALIGN.RIGHT, space=0)
    foot = 'Tú Lệ Smart Breakfast · Hệ thống truy xuất nguồn gốc'
    if section:
        foot += ' · ' + section
    text(s, ML, H - Inches(0.62), Inches(8.5), Inches(0.3), foot,
         size=10, color=MUTED, space=0)
    return s


def divider(number, title_text, subtitle):
    _n[0] += 1
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, W, H, fill=INK)
    box(s, 0, H - Inches(0.55), W, Inches(0.55), fill=ORANGE)
    text(s, ML, Inches(2.35), Inches(1.6), Inches(1.0), number,
         size=64, color=ORANGE, bold=True, space=0, line=1.0)
    text(s, ML, Inches(3.6), Inches(9.5), Inches(0.9), title_text,
         size=38, color=WHITE, bold=True, space=0)
    text(s, ML, Inches(4.65), Inches(9.5), Inches(0.9), subtitle,
         size=15, color=RGBColor(0xB9, 0xCD, 0xC1), space=0)
    return s


def bullets(s, items, left=ML, top=TOP, width=CW, size=15, gap=Inches(0.52),
            body_size=12.5):
    y = top
    for item in items:
        head, body = item if isinstance(item, tuple) else (item, None)
        box(s, left, y + Inches(0.09), Inches(0.13), Inches(0.13), fill=ORANGE)
        text(s, left + Inches(0.34), y, width - Inches(0.34), Inches(0.34),
             head, size=size, color=INK, bold=body is not None, space=2)
        if body:
            text(s, left + Inches(0.34), y + Inches(0.31),
                 width - Inches(0.34), Inches(0.4), body,
                 size=body_size, color=MUTED, space=0)
            y += gap + Inches(0.2)
        else:
            y += gap
    return y


def numbered_rows(s, items, left, top, width, row_h):
    """Danh sách đánh số, ngăn nhau bằng một đường kẻ mảnh."""
    for i, (head, body) in enumerate(items):
        y = top + row_h * i
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, int(left), int(y),
                                  Inches(0.34), Inches(0.34))
        circ.fill.background()
        circ.line.color.rgb = ORANGE
        circ.line.width = Pt(1.1)
        circ.shadow.inherit = False
        _txt(circ, str(i + 1), size=12, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER, space=0)
        tx = left + Inches(0.55)
        tw = width - Inches(0.55)
        text(s, tx, y + Inches(0.02), tw, Inches(0.3), head,
             size=15, color=INK, bold=True, space=3)
        text(s, tx, y + Inches(0.36), tw, row_h - Inches(0.5), body,
             size=12.5, color=MUTED, space=0)
        if i < len(items) - 1:
            box(s, tx, y + row_h - Inches(0.18), tw, Pt(0.9), fill=HAIR)


def table(s, rows, left=ML, top=TOP, width=CW, col_w=None, head_size=11.5,
          body_size=11.5, row_h=Inches(0.4)):
    nrow, ncol = len(rows), len(rows[0])
    gf = s.shapes.add_table(nrow, ncol, int(left), int(top), int(width),
                            int(row_h * nrow))
    t = gf.table
    t.first_row = False
    t.horz_banding = False
    if col_w:
        total = float(sum(col_w))
        for i, part in enumerate(col_w):
            t.columns[i].width = Emu(int(width * part / total))
    for r, row in enumerate(rows):
        t.rows[r].height = int(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                INK if r == 0 else (WHITE if r % 2 else BG2))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.space_after = Pt(0)
            p.line_spacing = 1.15
            run = p.add_run()
            run.text = str(val)
            f = run.font
            f.name = FONT
            f.size = Pt(head_size if r == 0 else body_size)
            f.bold = r == 0
            f.color.rgb = WHITE if r == 0 else INK
    return t


def cards(s, items, top=TOP, height=Inches(1.55), cols=3, gap=Inches(0.28),
          numbered=False, tint=WHITE, left=ML, width=CW):
    w = int((width - gap * (cols - 1)) / cols)
    for i, (head, body) in enumerate(items):
        r, c = divmod(i, cols)
        x = left + (w + gap) * c
        y = top + (height + gap) * r
        box(s, x, y, w, height, fill=tint, edge=LINE, radius=True)
        pad = Inches(0.24)
        ty = y + pad
        if numbered:
            circ = s.shapes.add_shape(MSO_SHAPE.OVAL, int(x + pad), int(ty),
                                      Inches(0.32), Inches(0.32))
            circ.fill.solid()
            circ.fill.fore_color.rgb = INK
            circ.line.fill.background()
            circ.shadow.inherit = False
            _txt(circ, str(i + 1), size=12, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, space=0)
            ty += Inches(0.44)
        text(s, x + pad, ty, w - pad * 2, Inches(0.3), head,
             size=13.5, color=INK, bold=True, space=3)
        text(s, x + pad, ty + Inches(0.32), w - pad * 2,
             height - (ty - y) - Inches(0.4), body,
             size=11.5, color=MUTED, space=0, line=1.2)


def stat(s, items, top=TOP, height=Inches(1.5)):
    n = len(items)
    gap = Inches(0.28)
    w = int((CW - gap * (n - 1)) / n)
    for i, (big, lab) in enumerate(items):
        x = ML + (w + gap) * i
        box(s, x, top, w, height, fill=WHITE, edge=LINE, radius=True)
        text(s, x + Inches(0.22), top + Inches(0.22), w - Inches(0.44),
             Inches(0.6), big, size=26, color=INK, bold=True, space=0)
        text(s, x + Inches(0.22), top + Inches(0.88), w - Inches(0.44),
             Inches(0.5), lab, size=11.5, color=MUTED, space=0, line=1.15)


def picture(s, name, left, top, width, height, caption=None):
    path = os.path.join(SHOT, name)
    ok = False
    if os.path.exists(path):
        try:
            from PIL import Image
            with Image.open(path) as im:
                iw, ih = im.size
            scale = min(float(width) / iw, float(height) / ih)
            pw, ph = int(iw * scale), int(ih * scale)
            s.shapes.add_picture(path, int(left + (width - pw) / 2),
                                 int(top + (height - ph) / 2), pw, ph)
            ok = True
        except ImportError:
            s.shapes.add_picture(path, int(left), int(top), int(width),
                                 int(height))
            ok = True
    if not ok:
        box(s, left, top, width, height, fill=WHITE, edge=LINE, dash=True,
            radius=True)
        text(s, left + Inches(0.2), top + height / 2 - Inches(0.3),
             width - Inches(0.4), Inches(0.6),
             ['Ảnh màn hình', 'tai-lieu/anh/' + name],
             size=11, color=MUTED, align=PP_ALIGN.CENTER, space=2)
    if caption:
        text(s, left, top + height + Inches(0.06), width, Inches(0.3),
             caption, size=10.5, color=MUTED, align=PP_ALIGN.CENTER, space=0)


def flow(s, steps, top, height=Inches(0.95), size=12):
    n = len(steps)
    arrow = Inches(0.34)
    w = int((CW - arrow * (n - 1)) / n)
    for i, lab in enumerate(steps):
        x = ML + (w + arrow) * i
        b = box(s, x, top, w, height, fill=WHITE, edge=LINE, radius=True)
        b.text_frame.margin_left = b.text_frame.margin_right = Inches(0.1)
        _txt(b, lab, size=size, color=INK, bold=True,
             align=PP_ALIGN.CENTER, space=0, line=1.15)
        b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        if i < n - 1:
            a = s.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, int(x + w + Inches(0.06)),
                int(top + height / 2 - Inches(0.09)),
                Inches(0.22), Inches(0.18))
            a.fill.solid()
            a.fill.fore_color.rgb = ORANGE
            a.line.fill.background()
            a.shadow.inherit = False


TECH = 'Phần A · Tài liệu kỹ thuật'
USE = 'Phần B · Hướng dẫn sử dụng'

# ------------------------------------------------------------------ slide 01
_n[0] += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=INK)
box(s, 0, 0, Inches(0.16), H, fill=ORANGE)
_px = W - ML - Inches(4.6)
for _img, _h, _y in ((os.path.join(ASSET, 'tule-products.jpg'), Inches(2.70),
                      Inches(1.38)),
                     (os.path.join(ASSET, 'tay-bac.jpg'), Inches(1.74),
                      Inches(4.38))):
    if os.path.exists(_img):
        s.shapes.add_picture(_img, int(_px), int(_y), int(Inches(4.6)),
                             int(_h))
text(s, ML, Inches(1.45), Inches(6.6), Inches(0.4),
     'CUỘC THI KHOA HỌC KỸ THUẬT', size=12, color=ORANGE, bold=True, space=0)
text(s, ML, Inches(2.0), Inches(6.6), Inches(2.2),
     'Hệ thống truy xuất nguồn gốc\nTú Lệ Smart Breakfast',
     size=34, color=WHITE, bold=True, space=0, line=1.15)
text(s, ML, Inches(3.7), Inches(6.4), Inches(1.2),
     'Ứng dụng mã QR và blockchain để chứng minh tính toàn vẹn của hồ sơ '
     'truy xuất, từ vùng nguyên liệu tới gói thành phẩm',
     size=14.5, color=PALE, space=0, line=1.35)
box(s, ML, Inches(4.85), Inches(2.9), Pt(1.4), fill=RGBColor(0x3C, 0x5C, 0x4C))
text(s, ML, Inches(5.15), Inches(6.4), Inches(1.4),
     [('Tài liệu kỹ thuật và hướng dẫn sử dụng',
       {'size': 13, 'bold': True, 'color': WHITE}),
      ('tule-trace.sontm.workers.dev  ·  VBSN Besu, chainId 84001',
       {'size': 12, 'color': RGBColor(0x9F, 0xB8, 0xA9)})], space=4)

# ------------------------------------------------------------------ slide 02
s = slide('Một tem truy xuất nguồn gốc hiện nay chứng minh được gì?')
LW = Inches(6.5)
numbered_rows(s, [
    ('Dữ liệu nằm trên máy chủ của người bán',
     'Người bán vừa lập hồ sơ vừa giữ hồ sơ. Sửa nội dung sau khi in bao bì '
     'không để lại dấu vết nào cho người mua kiểm tra.'),
    ('Nội dung không gắn với lô hàng cụ thể',
     'Phần lớn tem chỉ dẫn tới một trang giới thiệu thương hiệu, không có dữ '
     'liệu của đúng lô hàng đang nằm trong tay người mua.'),
    ('Không có mốc thời gian độc lập',
     'Không có căn cứ nào để xác định hồ sơ được lập trước hay sau khi hàng '
     'đã ra thị trường.'),
], ML, TOP, LW, Inches(1.42))
_rx = ML + LW + Inches(0.45)
_rw = CW - LW - Inches(0.45)
box(s, _rx, TOP, _rw, Inches(4.26), fill=GREEN_SOFT, edge=GREEN_EDGE,
    radius=True)
label(s, _rx + Inches(0.32), TOP + Inches(0.3), _rw - Inches(0.64),
      'Câu hỏi nghiên cứu')
text(s, _rx + Inches(0.32), TOP + Inches(0.75), _rw - Inches(0.64),
     Inches(3.2),
     'Làm thế nào để người tiêu dùng tự kiểm tra được rằng hồ sơ của một lô '
     'hàng không bị sửa sau ngày công bố, mà không phải tin vào lời của nhà '
     'sản xuất và không cần kiến thức kỹ thuật?',
     size=17, color=INK, space=0, line=1.4)

# ------------------------------------------------------------------ slide 03
s = slide('Mục tiêu và phạm vi đề tài')
bullets(s, [
    ('Số hoá trọn quy trình sản xuất thực tế',
     'Bốn nhánh nguyên liệu song song (cốm Tú Lệ, lạc đỏ Lục Yên, chuối tiêu '
     'xanh, khoai môn Lục Yên) với 22 công đoạn, cộng phối trộn và đóng gói.'),
    ('Cố định hồ sơ bằng mã băm và neo lên blockchain công khai',
     'Mỗi lần công bố sinh một phiên bản mới kèm mã băm SHA-256; mã băm được '
     'ghi thành một giao dịch trên VBSN Besu.'),
    ('Người tiêu dùng kiểm chứng bằng một lần quét mã QR',
     'Không cài ứng dụng, không tạo tài khoản, không cần ví điện tử.'),
    ('Người vận hành thao tác trên một màn duy nhất cho mỗi lô',
     'Hệ thống tạo sẵn toàn bộ công đoạn theo sơ đồ quy trình; phần việc chính '
     'còn lại là tải ảnh thực tế lên.'),
], size=16)

# ------------------------------------------------------------------ slide 04
s = slide('Toàn bộ hệ thống trong một hình')
flow(s, ['Người vận hành\nnhập dữ liệu lô',
         'Hệ thống chuẩn hoá\nvà băm SHA-256',
         'Ghi mã băm lên\nVBSN Besu',
         'In mã QR\nlên bao bì',
         'Người mua quét\nvà đối chiếu'], Inches(2.15), Inches(1.15))
text(s, ML, Inches(3.63), CW, Inches(0.4),
     'Ba tính chất kỹ thuật tạo nên giá trị chứng minh',
     size=13.5, color=INK, bold=True, space=0)
cards(s, [
    ('Bất biến theo phiên bản',
     'Công bố lần sau không ghi đè lần trước. Toàn bộ nội dung cũ và giao dịch '
     'tương ứng vẫn tra cứu lại được.'),
    ('Mốc thời gian độc lập',
     'Thời điểm và mã băm nằm trong một khối đã xác nhận của blockchain, '
     'không do hệ thống tự khai báo.'),
    ('Kiểm chứng không cần tin hệ thống',
     'Bất kỳ ai cũng đọc được sự kiện trực tiếp từ blockchain và tự băm lại '
     'nội dung để so sánh.'),
], top=Inches(4.15), height=Inches(1.75))

# ------------------------------------------------------------------ slide 05
divider('A', 'Tài liệu kỹ thuật',
        'Kiến trúc, mô hình dữ liệu, cơ chế toàn vẹn, tiêu chuẩn, kiểm thử')

# ------------------------------------------------------------------ slide 06
s = slide('Kiến trúc hệ thống', TECH)
col = int((CW - Inches(0.5) * 2) / 3)
x0 = ML
box(s, x0, TOP, col, Inches(1.95), fill=WHITE, edge=LINE, radius=True)
label(s, x0 + Inches(0.24), TOP + Inches(0.22), col - Inches(0.48),
      'Người dùng')
text(s, x0 + Inches(0.24), TOP + Inches(0.62), col - Inches(0.48), Inches(1.0),
     [('Người tiêu dùng: điện thoại, vào bằng mã QR', {'size': 12}),
      ('Người vận hành: máy tính, mở màn /admin', {'size': 12})],
     space=6, color=INK)
x1 = x0 + col + Inches(0.5)
box(s, x1, TOP - Inches(0.12), col, Inches(4.15), fill=INK, radius=True)
label(s, x1 + Inches(0.24), TOP + Inches(0.04), col - Inches(0.48),
      'Cloudflare')
inner = [
    'Worker (Hono): định tuyến, xác thực, nghiệp vụ',
    'Static Assets: gói build Flutter Web',
    'D1 (SQLite): lô, nguyên liệu, công đoạn, bản công bố',
    'R2: ảnh, video, hồ sơ kiểm nghiệm',
    'Cron Trigger: quét hàng chờ mỗi 10 phút',
]
for i, lab in enumerate(inner):
    yy = TOP + Inches(0.45) + Inches(0.7) * i
    b = box(s, x1 + Inches(0.24), yy, col - Inches(0.48), Inches(0.58),
            fill=INK_SOFT, radius=True)
    b.text_frame.margin_left = b.text_frame.margin_right = Inches(0.14)
    _txt(b, lab, size=11.5, color=WHITE, space=0, line=1.1)
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
x2 = x1 + col + Inches(0.5)
box(s, x2, TOP, col, Inches(1.95), fill=WHITE, edge=LINE, radius=True)
label(s, x2 + Inches(0.24), TOP + Inches(0.22), col - Inches(0.48),
      'Dịch vụ ngoài')
text(s, x2 + Inches(0.24), TOP + Inches(0.62), col - Inches(0.48), Inches(1.0),
     [('VBSN Besu: chainId 84001, hợp đồng TuleTrace', {'size': 12}),
      ('OpenFreeMap: dữ liệu nền bản đồ', {'size': 12})],
     space=6, color=INK)
box(s, x2, TOP + Inches(2.1), col, Inches(1.93), fill=GREEN_SOFT,
    edge=GREEN_EDGE, radius=True)
label(s, x2 + Inches(0.24), TOP + Inches(2.28), col - Inches(0.48),
      'Đặc điểm triển khai')
text(s, x2 + Inches(0.24), TOP + Inches(2.64), col - Inches(0.48), Inches(1.3),
     [('Một Worker phục vụ cả giao diện và API.', {'size': 11.5}),
      ('Không có máy chủ thường trực.', {'size': 11.5}),
      ('Không tính chi phí theo giờ.', {'size': 11.5}),
      ('Khoá ký nằm trong secret của Worker.', {'size': 11.5})],
     space=6, color=INK, line=1.15)

# ------------------------------------------------------------------ slide 07
s = slide('Công nghệ sử dụng', TECH)
table(s, [
    ['Thành phần', 'Công nghệ', 'Trách nhiệm'],
    ['Giao diện web', 'Flutter Web 3.44 (CanvasKit), go_router',
     'Trang tra cứu công khai và màn quản trị'],
    ['Lớp API', 'Cloudflare Workers, Hono 4',
     'Định tuyến, xác thực, kiểm tra đầu vào, điều phối nghiệp vụ'],
    ['Cơ sở dữ liệu', 'Cloudflare D1 (SQLite)',
     'Lô, nguyên liệu, công đoạn, media, bản công bố, hàng chờ neo chuỗi'],
    ['Kho tệp', 'Cloudflare R2',
     'Ảnh, video, hồ sơ kiểm nghiệm; khoá đặt theo nội dung tệp'],
    ['Bản đồ', 'MapLibre GL JS 6.4.1, OpenFreeMap Liberty',
     'Hiển thị vùng nguyên liệu và tuyến vận chuyển'],
    ['Tích hợp chuỗi', 'viem 2.x', 'Ký và gửi giao dịch neo mã băm'],
    ['Blockchain', 'VBSN Besu (EVM), chainId 84001',
     'Lưu mã băm và mốc thời gian công bố'],
], col_w=[3, 5, 7], row_h=Inches(0.55), body_size=12)

# ------------------------------------------------------------------ slide 08
s = slide('Mô hình dữ liệu', TECH)
table(s, [
    ['Bảng', 'Nội dung', 'Ràng buộc'],
    ['product_batches',
     'Lô thành phẩm: mã truy xuất, ngày sản xuất, hạn dùng, nơi sản xuất, '
     'toạ độ, trạng thái', 'Mã lô duy nhất toàn hệ thống'],
    ['ingredient_batches',
     'Lô nguyên liệu: nhà cung cấp, ngày thu hoạch, ngày nhập kho, vùng trồng '
     'đã khoanh', 'Xoá lô kéo theo dữ liệu phụ thuộc'],
    ['process_events', 'Công đoạn, có thứ tự hiển thị riêng',
     'CHECK: thuộc đúng một lô nguyên liệu hoặc một lô thành phẩm'],
    ['media_assets', 'Ảnh, video, hồ sơ kiểm nghiệm',
     'Khoá đặt theo nội dung: media/<sha256>.<ext>'],
    ['published_snapshots',
     'Bản công bố: phiên bản, toàn bộ nội dung, mã băm SHA-256, thời điểm',
     'Cặp (lô, phiên bản) là duy nhất'],
    ['blockchain_outbox',
     'Hàng chờ neo chuỗi: trạng thái, mã giao dịch, số khối',
     'Tối đa 5 lần thử gửi'],
], col_w=[3, 7, 5], row_h=Inches(0.62), body_size=12)
text(s, ML, Inches(6.3), CW, Inches(0.4),
     'Sơ đồ quan hệ đầy đủ giữa sáu bảng nằm ở KIEN_TRUC.md mục 4.',
     size=11.5, color=MUTED, space=0)

# ------------------------------------------------------------------ slide 09
s = slide('Cơ chế toàn vẹn: chuẩn hoá rồi mới băm', TECH)
text(s, ML, TOP - Inches(0.1), CW, Inches(0.4),
     'Cùng một nội dung phải cho ra cùng một mã băm, bất kể thứ tự truy vấn '
     'hay phiên bản hệ thống. Quy tắc chuẩn hoá:',
     size=13, color=MUTED, space=0)
bullets(s, [
    'Khoá JSON sắp xếp theo thứ tự từ điển ở mọi cấp',
    'Loại bỏ định danh nội bộ, dấu thời gian hệ thống và trường phụ thuộc '
    'hạ tầng',
    'Trường không có dữ liệu bị loại khỏi payload, không ghi giá trị rỗng',
    'Tệp đính kèm tham chiếu bằng khoá nội dung media/<sha256>.<ext>',
    'Hàm băm SHA-256, kết quả hexa chữ thường 64 ký tự',
], top=TOP + Inches(0.5), size=14, gap=Inches(0.5))
box(s, ML, Inches(5.3), CW, Inches(1.2), fill=GREEN_SOFT, edge=GREEN_EDGE,
    radius=True)
label(s, ML + Inches(0.3), Inches(5.5), CW - Inches(0.6),
      'Vai trò của quy tắc thứ ba')
text(s, ML + Inches(0.3), Inches(5.8), CW - Inches(0.6), Inches(0.6),
     'Loại trường rỗng khỏi payload là điều kiện tương thích ngược: bổ sung '
     'một trường mới vào hệ thống không làm thay đổi mã băm của những lô đã '
     'công bố trước đó và chưa dùng trường ấy.',
     size=13, color=INK, space=0)

# ------------------------------------------------------------------ slide 10
s = slide('Luồng công bố và neo chuỗi', TECH)
cards(s, [
    ('Gom dữ liệu', 'Truy vấn toàn bộ dữ liệu của lô'),
    ('Chuẩn hoá và băm', 'Sinh JSON chuẩn hoá, tính SHA-256'),
    ('Lưu phiên bản', 'INSERT published_snapshots, version = max + 1'),
    ('Vào hàng chờ', 'INSERT blockchain_outbox, PENDING_NETWORK'),
    ('Gửi giao dịch', 'anchorBatch(...) qua viem, chạy nền bằng waitUntil'),
    ('Xác nhận', 'Ghi tx_hash, block_number, trạng thái CONFIRMED'),
], top=TOP, height=Inches(1.5), cols=3, numbered=True)
box(s, ML, Inches(5.4), CW, Inches(1.05), fill=BG2, edge=LINE, radius=True)
label(s, ML + Inches(0.28), Inches(5.58), CW - Inches(0.56),
      'Cơ chế phục hồi')
text(s, ML + Inches(0.28), Inches(5.88), CW - Inches(0.56), Inches(0.45),
     'Cron chạy mỗi 10 phút quét lại bản ghi PENDING_NETWORK và FAILED có số '
     'lần thử dưới 5, xử lý trường hợp mạng lỗi ngay tại thời điểm công bố. '
     'Dữ liệu không mất, chỉ chậm có giao dịch.',
     size=12, color=INK, space=0)

# ------------------------------------------------------------------ slide 11
s = slide('Hợp đồng thông minh TuleTrace', TECH)
text(s, ML, TOP - Inches(0.1), CW, Inches(0.4),
     'Hợp đồng chỉ phát sự kiện, không lưu nội dung hồ sơ. Dữ liệu cá nhân và '
     'dữ liệu chi phí không được đưa lên chuỗi.',
     size=13, color=MUTED, space=0)
table(s, [
    ['Loại', 'Chữ ký', 'Điều kiện gọi'],
    ['Sự kiện', 'BatchAnchored(bytes32 hash, address by, string code, '
     'uint32 version, bytes32[] inputs, uint64 at)', 'Không áp dụng'],
    ['Sự kiện', 'IngredientAnchored(bytes32 hash, address by, string ref, '
     'uint64 at)', 'Không áp dụng'],
    ['Hàm', 'anchorBatch(bytes32, string, uint32, bytes32[])', 'Vai trò BRAND'],
    ['Hàm', 'anchorIngredient(bytes32, string)', 'Vai trò khác NONE'],
    ['Hàm', 'setActor(address, uint8, string)', 'Chủ sở hữu hợp đồng'],
], top=TOP + Inches(0.45), col_w=[2, 8, 3], row_h=Inches(0.5), body_size=11.5)
stat(s, [('29.273 gas', 'cho một lần neo mã băm'),
         ('0,0000000293 VNX', 'chi phí quy đổi mỗi giao dịch'),
         ('0xcd6811f9…d37ffd3ca', 'địa chỉ hợp đồng đang vận hành')],
     top=Inches(5.55), height=Inches(1.1))

# ------------------------------------------------------------------ slide 12
s = slide('Hai phép kiểm chứng, hai ý nghĩa khác nhau', TECH)
box(s, ML, TOP, HALF, Inches(2.95), fill=WHITE, edge=LINE, radius=True)
label(s, ML + Inches(0.3), TOP + Inches(0.26), HALF - Inches(0.6),
      'Kiểm chứng công khai')
text(s, ML + Inches(0.3), TOP + Inches(0.66), HALF - Inches(0.6), Inches(2.1),
     [('So sánh nội dung đang phục vụ với mã băm đã lưu và đã neo lên chuỗi.',
       {'size': 13}),
      ('Trang quét QR phục vụ nội dung của bản công bố, không phải dữ liệu '
       'hiện hành trong màn quản trị.', {'size': 13}),
      ('Kết quả DỮ LIỆU SAI LỆCH nghĩa là dữ liệu công bố đã bị can thiệp ở '
       'tầng lưu trữ.', {'size': 13, 'color': INK, 'bold': True})],
     space=9, color=MUTED, line=1.25)
x = ML + HALF + Inches(0.4)
box(s, x, TOP, HALF, Inches(2.95), fill=WHITE, edge=LINE, radius=True)
label(s, x + Inches(0.3), TOP + Inches(0.26), HALF - Inches(0.6),
      'Kiểm chứng quản trị')
text(s, x + Inches(0.3), TOP + Inches(0.66), HALF - Inches(0.6), Inches(2.1),
     [('So sánh dữ liệu hiện hành với bản công bố gần nhất.', {'size': 13}),
      ('Dùng để người vận hành biết mình đã sửa những gì mà chưa công bố.',
       {'size': 13}),
      ('Kết quả CÓ THAY ĐỔI CHƯA LƯU là trạng thái làm việc bình thường, '
       'không phải cảnh báo.', {'size': 13, 'color': INK, 'bold': True})],
     space=9, color=MUTED, line=1.25)
box(s, ML, Inches(5.25), CW, Inches(1.1), fill=INK, radius=True)
label(s, ML + Inches(0.3), Inches(5.45), CW - Inches(0.6), 'Hệ quả thiết kế')
text(s, ML + Inches(0.3), Inches(5.75), CW - Inches(0.6), Inches(0.5),
     'Sửa dữ liệu trong màn quản trị không làm thay đổi ngay nội dung người '
     'mua nhìn thấy. Nội dung mới chỉ được phục vụ sau khi công bố một phiên '
     'bản mới, và phiên bản đó có mã băm riêng.',
     size=13, color=WHITE, space=0)

# ------------------------------------------------------------------ slide 13
s = slide('Ràng buộc thiết kế', TECH)
table(s, [
    ['Mã', 'Ràng buộc'],
    ['RB-01', 'Khi API không phản hồi, giao diện hiển thị trạng thái lỗi; hệ '
     'thống không thay thế bằng dữ liệu mẫu'],
    ['RB-02', 'Nhãn giao diện phản ánh đúng trạng thái kỹ thuật: trạng thái đã '
     'lưu lên blockchain chỉ hiển thị khi tồn tại giao dịch đã xác nhận'],
    ['RB-03', 'Trường không có dữ liệu không được đưa vào payload công bố'],
    ['RB-04', 'Không xoá tệp khỏi kho khi còn bản công bố tham chiếu tới nó'],
    ['RB-05', 'Bán kính vùng nguyên liệu do người dùng khai không được vẽ '
     'thành hình tròn trên bản đồ; bản đồ chỉ tô vùng được khoanh trực tiếp'],
    ['RB-06', 'Dữ liệu chi phí, giá vốn và lợi nhuận không được đưa vào '
     'payload công bố'],
    ['RB-07', 'Công bố không ghi đè bản trước; mỗi lần công bố tạo một phiên '
     'bản mới'],
], col_w=[1, 11], row_h=Inches(0.56), body_size=12)

# ------------------------------------------------------------------ slide 14
s = slide('Đối chiếu tiêu chuẩn truy xuất nguồn gốc', TECH)
table(s, [
    ['Văn bản, tiêu chuẩn', 'Yêu cầu chính', 'Mức đáp ứng'],
    ['Thông tư 02/2024/TT-BKHCN',
     'Dữ liệu truy xuất có mã định danh, không sửa được sau công bố',
     'Đáp ứng'],
    ['TCVN 13274:2020', 'Định dạng mã truy xuất nguồn gốc',
     'Dùng mã tự đặt, chưa dùng mã GS1'],
    ['TCVN 13275:2020', 'Định dạng vật mang dữ liệu (mã QR)',
     'Mã QR chứa URL thuần, chưa theo GS1 Digital Link'],
    ['TCVN 12850:2019', 'Yêu cầu với dữ liệu và vật mang dữ liệu',
     'Đáp ứng phần dữ liệu'],
    ['GS1 EPCIS', 'Mô hình sự kiện biến đổi (TransformationEvent)',
     'Đáp ứng về cấu trúc, chưa dùng mã GS1'],
    ['ISO 22005', 'Truy xuất trong chuỗi thực phẩm',
     'Đáp ứng phần định danh, thiếu cân bằng khối lượng'],
], col_w=[4, 6, 4], row_h=Inches(0.56), body_size=12)
box(s, ML, Inches(5.95), CW, Inches(0.82), fill=BG2, edge=LINE, radius=True)
text(s, ML + Inches(0.28), Inches(6.14), CW - Inches(0.56), Inches(0.5),
     'Ba dòng liên quan tới mã GS1 phụ thuộc một thủ tục hành chính, không '
     'phải hạng mục lập trình. Trang sau nói rõ ai làm phần nào.',
     size=13, color=INK, bold=True, space=0)

# ------------------------------------------------------------------ slide 15
s = slide('Vì sao hệ thống chưa dùng mã GS1', TECH)
text(s, ML, TOP - Inches(0.1), CW, Inches(0.5),
     'GTIN (mã sản phẩm) và GLN (mã địa điểm) không phải mã tự đặt được. '
     'Chúng sinh ra từ mã doanh nghiệp GS1 mà GS1 Việt Nam cấp cho từng doanh '
     'nghiệp. Tự đặt một dãy số sẽ trùng vào mã của doanh nghiệp khác.',
     size=13, color=MUTED, space=0, line=1.35)
box(s, ML, TOP + Inches(0.68), HALF, Inches(3.05), fill=WHITE, edge=LINE,
    radius=True)
label(s, ML + Inches(0.3), TOP + Inches(0.9), HALF - Inches(0.6),
      'Phần hệ thống · khoảng nửa ngày công')
text(s, ML + Inches(0.3), TOP + Inches(1.28), HALF - Inches(0.6), Inches(2.2),
     [('Thêm trường gtin cho lô và gln cho nơi sản xuất.', {'size': 12.5}),
      ('Đổi nội dung mã QR sang GS1 Digital Link, dạng '
       'https://ten-mien/01/{GTIN}/10/{số lô}.', {'size': 12.5}),
      ('Giữ song song đường dẫn /t/{mã lô} để tem đã in vẫn mở được.',
       {'size': 12.5}),
      ('Kiến trúc không đổi: vẫn là một URL mở ra một trang.',
       {'size': 12.5})], space=7, color=MUTED, line=1.25)
x = ML + HALF + Inches(0.4)
box(s, x, TOP + Inches(0.68), HALF, Inches(3.05), fill=INK, radius=True)
label(s, x + Inches(0.3), TOP + Inches(0.9), HALF - Inches(0.6),
      'Phần doanh nghiệp · thủ tục với GS1 Việt Nam')
text(s, x + Inches(0.3), TOP + Inches(1.28), HALF - Inches(0.6), Inches(2.2),
     [('Hồ sơ: giấy chứng nhận đăng ký kinh doanh, bản đăng ký sử dụng mã số '
       'mã vạch, bảng danh mục sản phẩm dùng GTIN.', {'size': 12.5}),
      ('Phí cấp mã doanh nghiệp GS1 là 1.000.000 đồng, mã địa điểm GLN là '
       '300.000 đồng, cộng phí duy trì hằng năm theo loại mã.',
       {'size': 12.5}),
      ('Nộp trực tuyến qua cổng VNPC, hồ sơ giấy tại Trung tâm Mã số Mã vạch '
       'Quốc gia.', {'size': 12.5})],
     space=8, color=RGBColor(0xD5, 0xE2, 0xD9), line=1.25)
box(s, ML, Inches(5.75), CW, Inches(0.9), fill=GREEN_SOFT, edge=GREEN_EDGE,
    radius=True)
text(s, ML + Inches(0.3), Inches(5.95), CW - Inches(0.6), Inches(0.5),
     'Lưu ý thời điểm: nội dung mã QR nằm cứng trong bao bì đã in. Nếu có ý '
     'định nối Cổng truy xuất nguồn gốc quốc gia hoặc đưa hàng vào hệ thống '
     'bán lẻ, nên đăng ký GS1 trước khi in hàng loạt.',
     size=13, color=INK, bold=True, space=0)

# ------------------------------------------------------------------ slide 16
s = slide('Kiểm thử và số liệu vận hành', TECH)
table(s, [
    ['Hạng mục', 'Công cụ', 'Phạm vi'],
    ['Phân tích tĩnh', 'flutter analyze', 'Toàn bộ mã Dart'],
    ['Kiểm thử giao diện', 'flutter test',
     '39 trường hợp: bộ phân tích dữ liệu, bố cục ở 360/768/1440 px, bản đồ, '
     'luồng quản trị'],
    ['Kiểm kiểu backend', 'npx tsc --noEmit', 'Toàn bộ mã TypeScript'],
    ['Kiểm thử tích hợp', 'node scripts/smoke-test.js',
     'Xác thực, mã lỗi, tải tệp, HTTP Range, công bố, hàng chờ neo chuỗi'],
    ['Kiểm chứng độc lập', 'node scripts/chain-verify.mjs',
     'Đọc sự kiện trực tiếp từ blockchain, không qua máy chủ hệ thống'],
], col_w=[3, 4, 8], row_h=Inches(0.56), body_size=12)
stat(s, [('≈ 4,4 MB', 'dung lượng lần tải đầu trên di động'),
         ('≈ 4 giây', 'thời gian tới trạng thái networkidle'),
         ('10 phút', 'chu kỳ quét lại hàng chờ neo chuỗi'),
         ('30 MB', 'giới hạn mỗi tệp tải lên')],
     top=Inches(5.35), height=Inches(1.25))

# ------------------------------------------------------------------ slide 17
divider('B', 'Hướng dẫn sử dụng',
        'Từ lúc quét mã QR tới lúc công bố lô và in bao bì')

# ------------------------------------------------------------------ slide 18
s = slide('Hai giao diện của hệ thống', USE)
box(s, ML, TOP, HALF, Inches(1.2), fill=WHITE, edge=LINE, radius=True)
label(s, ML + Inches(0.3), TOP + Inches(0.24), HALF - Inches(0.6),
      'Trang khách')
text(s, ML + Inches(0.3), TOP + Inches(0.6), HALF - Inches(0.6), Inches(0.5),
     'tule-trace.sontm.workers.dev, mở công khai, không cần đăng nhập',
     size=13, color=INK, space=0)
picture(s, '02-ho-so-lo.png', ML, TOP + Inches(1.45), HALF, Inches(2.75),
        'Hồ sơ một lô: bản đồ vùng nguyên liệu bên trái, hành trình bên phải')
x = ML + HALF + Inches(0.4)
box(s, x, TOP, HALF, Inches(1.2), fill=WHITE, edge=LINE, radius=True)
label(s, x + Inches(0.3), TOP + Inches(0.24), HALF - Inches(0.6),
      'Màn quản trị')
text(s, x + Inches(0.3), TOP + Inches(0.6), HALF - Inches(0.6), Inches(0.5),
     'Đường dẫn /admin, nội bộ, mở bằng token quản trị, trình duyệt nhớ token',
     size=13, color=INK, space=0)
picture(s, '09-man-lam-viec.png', x, TOP + Inches(1.45), HALF, Inches(2.75),
        'Màn làm việc của một lô: mọi thao tác nằm trên một trang')

# ------------------------------------------------------------------ slide 19
s = slide('Người tiêu dùng: ba bước', USE)
cards(s, [
    ('Quét mã QR trên bao bì',
     'Camera điện thoại mở thẳng hồ sơ lô. Không quét được thì nhập mã in '
     'dưới mã QR, hoặc bấm nút quét camera trong trang.'),
    ('Xem hành trình của lô',
     'Bốn vùng nguyên liệu trên bản đồ, từng công đoạn kèm ảnh thực tế, nhà '
     'cung cấp và ngày thu hoạch.'),
    ('Đối chiếu bằng chứng',
     'Dải cuối trang hiện phiên bản, thời điểm công bố, mã giao dịch và nút '
     'mở giao dịch trên explorer.'),
], height=Inches(1.95), numbered=True)
picture(s, '03-nguyen-lieu.png', ML, Inches(4.4), HALF, Inches(2.2),
        'Hồ sơ một nguyên liệu')
picture(s, '04-kiem-chung.png', ML + HALF + Inches(0.4), Inches(4.4), HALF,
        Inches(2.2), 'Dải kiểm chứng cuối trang')

# ------------------------------------------------------------------ slide 20
s = slide('Người vận hành: tạo một lô mới', USE)
table(s, [
    ['Trường', 'Ghi chú'],
    ['Tên lô', 'Nội dung người mua nhìn thấy đầu tiên khi quét mã QR'],
    ['Mã truy xuất', 'In lên bao bì và nằm trong đường dẫn QR. Chữ in hoa, số, '
     'dấu gạch ngang. Trùng mã bị chặn'],
    ['Ngày sản xuất, hạn sử dụng', 'Chọn bằng lịch'],
    ['Nơi sản xuất', 'Mặc định: Trung tâm Phát triển và Giao dịch Công nghệ '
     'thành phố Hà Nội'],
    ['Tạo sẵn bốn nguyên liệu Tú Lệ',
     'Nên chọn. Hệ thống tạo luôn cốm, lạc, chuối, khoai môn kèm 22 công đoạn '
     'theo sơ đồ quy trình, cộng phối trộn và đóng gói'],
], col_w=[3, 9], row_h=Inches(0.58), body_size=12)
box(s, ML, Inches(5.5), CW, Inches(0.9), fill=GREEN_SOFT, edge=GREEN_EDGE,
    radius=True)
text(s, ML + Inches(0.3), Inches(5.72), CW - Inches(0.6), Inches(0.5),
     'Chọn ô cuối cùng thì phần việc còn lại của người vận hành gần như chỉ là '
     'tải ảnh thực tế lên từng công đoạn.',
     size=13.5, color=INK, bold=True, space=0)

# ------------------------------------------------------------------ slide 21
s = slide('Màn làm việc của một lô', USE)
cards(s, [
    ('Thông tin lô',
     'Tên, ngày sản xuất, hạn dùng, nơi sản xuất, toạ độ, mô tả ngắn. Nút lưu '
     'chỉ hoạt động khi có thay đổi thật.'),
    ('Ảnh của lô',
     'Ảnh giới thiệu, ảnh bìa và hồ sơ kiểm nghiệm. Mỗi tệp tối đa 30 MB, nén '
     'video xuống 720p trước khi tải.'),
    ('Nguyên liệu',
     'Nhà cung cấp, ngày thu hoạch, ngày nhập kho, vùng trồng khoanh trực tiếp '
     'trên bản đồ, và toàn bộ công đoạn kèm ảnh riêng.'),
    ('Công đoạn tại xưởng',
     'Phối trộn và đóng gói thuộc cả lô, không thuộc nguyên liệu nào, nên nằm '
     'ở khối riêng.'),
], cols=4, height=Inches(1.95))
picture(s, '12-nguyen-lieu-admin.png', ML, Inches(4.45), HALF, Inches(2.15),
        'Thẻ nguyên liệu mở ra')
picture(s, '13-khoanh-vung.png', ML + HALF + Inches(0.4), Inches(4.45), HALF,
        Inches(2.15), 'Khoanh vùng nguyên liệu trên bản đồ')

# ------------------------------------------------------------------ slide 22
s = slide('Hoàn thành lô: một lần bấm, bốn việc', USE)
flow(s, ['Gom dữ liệu\nthành một bản',
         'Băm SHA-256 và\nlưu phiên bản mới',
         'Gửi mã băm lên\nVBSN Besu',
         'Người mua quét QR\nthấy nội dung mới'], TOP, Inches(1.05))
text(s, ML, TOP + Inches(1.22), CW, Inches(0.4),
     'Người vận hành không phải xác nhận ví và không phải trả phí giao dịch: '
     'hệ thống ký bằng ví của thương hiệu.',
     size=13, color=MUTED, space=0)
box(s, ML, Inches(3.72), HALF, Inches(2.6), fill=WHITE, edge=LINE, radius=True)
label(s, ML + Inches(0.3), Inches(3.94), HALF - Inches(0.6),
      'Sau khi công bố, khối này hiện')
text(s, ML + Inches(0.3), Inches(4.32), HALF - Inches(0.6), Inches(1.85),
     [('Trạng thái dữ liệu khớp bản đã công bố, kèm mã băm.', {'size': 12.5}),
      ('Dòng đã lưu lên blockchain, kèm nút xem giao dịch.', {'size': 12.5}),
      ('Lịch sử các bản: bấm biểu tượng đồng hồ để xem lại nội dung một bản '
       'đúng như lúc chốt.', {'size': 12.5}),
      ('Mã QR của lô, tải được thành file PNG nền trắng 1024 px kèm mã lô, '
       'dùng thẳng cho bản in bao bì.', {'size': 12.5})],
     space=7, color=MUTED, line=1.25)
picture(s, '15-hoan-thanh-lo.png', ML + HALF + Inches(0.4), Inches(3.72), HALF,
        Inches(2.6), 'Khối hoàn thành lô trong màn quản trị')

# ------------------------------------------------------------------ slide 23
s = slide('Những trạng thái hay gặp', USE)
table(s, [
    ['Bạn thấy', 'Nghĩa là', 'Cần làm gì'],
    ['Có thay đổi chưa lưu', 'Vừa sửa dữ liệu; người mua vẫn đang xem bản cũ',
     'Bấm Lưu bản mới khi đã nhập xong'],
    ['CHỜ LƯU CHUỖI', 'Mã băm đã vào hàng chờ, chưa có giao dịch',
     'Đợi tối đa 10 phút, hoặc bấm Gửi lên chuỗi'],
    ['GỬI HỎNG', 'Giao dịch thất bại, thường do mạng hoặc ví hết VNX',
     'Kiểm số dư ví hiện dưới nút, rồi gửi lại'],
    ['DỮ LIỆU SAI LỆCH', 'Nội dung đang phục vụ không khớp mã băm của chính nó',
     'Hiếm gặp; báo kỹ thuật kiểm tra ngay'],
    ['ĐANG TẠO', 'Lô chưa công bố bản nào', 'Nhập xong thì bấm Hoàn thành lô'],
], col_w=[3, 6, 5], row_h=Inches(0.56), body_size=12)
box(s, ML, Inches(5.45), CW, Inches(0.95), fill=BG2, edge=LINE, radius=True)
text(s, ML + Inches(0.28), Inches(5.67), CW - Inches(0.56), Inches(0.5),
     'Sửa dữ liệu không xoá bản cũ: nội dung và giao dịch của bản trước vẫn '
     'còn nguyên, xem lại được trong Lịch sử. Ảnh đã gỡ khỏi lô vẫn được giữ '
     'nếu còn bản công bố nào trỏ tới nó.',
     size=12.5, color=INK, space=0, line=1.25)

# ------------------------------------------------------------------ slide 24
s = slide('Việc định kỳ và lưu ý trước khi in bao bì', USE)
box(s, ML, TOP, HALF, Inches(3.8), fill=WHITE, edge=LINE, radius=True)
label(s, ML + Inches(0.3), TOP + Inches(0.24), HALF - Inches(0.6),
      'Việc định kỳ')
text(s, ML + Inches(0.3), TOP + Inches(0.62), HALF - Inches(0.6), Inches(2.9),
     [('Xem tình trạng ví neo dữ liệu', {'size': 12.5, 'bold': True}),
      ('npm run chain:status', {'size': 11.5, 'color': MUTED}),
      ('Kiểm chứng độc lập một mã băm', {'size': 12.5, 'bold': True}),
      ('node scripts/chain-verify.mjs <contract> <sha256>',
       {'size': 11.5, 'color': MUTED}),
      ('Chép một lô từ máy thử lên production', {'size': 12.5, 'bold': True}),
      ('node scripts/copy-batch.mjs TL-2026-003 --to <url>',
       {'size': 11.5, 'color': MUTED}),
      ('Đưa công đoạn của một lô về đúng sơ đồ quy trình',
       {'size': 12.5, 'bold': True}),
      ('node scripts/apply-process.mjs TL-2026-003 --publish',
       {'size': 11.5, 'color': MUTED})], space=5, color=INK)
x = ML + HALF + Inches(0.4)
box(s, x, TOP, HALF, Inches(3.8), fill=INK, radius=True)
label(s, x + Inches(0.3), TOP + Inches(0.24), HALF - Inches(0.6),
      'Trước khi in bao bì')
text(s, x + Inches(0.3), TOP + Inches(0.62), HALF - Inches(0.6), Inches(2.9),
     [('Mã lô nằm trong đường dẫn QR. In rồi thì không đổi mã được nữa; đổi mã '
       'là mọi gói đã in trỏ vào một trang không tồn tại.', {'size': 12.5}),
      ('Đổi token quản trị sang chuỗi dài, khó đoán. Đây là thứ duy nhất chặn '
       'người lạ tạo lô giả dưới tên thương hiệu.', {'size': 12.5}),
      ('Kiểm tra bằng điện thoại thật: quét thử mã QR đã tải về, xem trang mở '
       'đúng lô và ảnh hiện đủ.', {'size': 12.5}),
      ('Ví neo dữ liệu phải còn VNX. Hết thì hàng chờ đứng lại, dữ liệu vẫn an '
       'toàn nhưng không có giao dịch mới.', {'size': 12.5})],
     space=9, color=RGBColor(0xD5, 0xE2, 0xD9), line=1.25)

# ------------------------------------------------------------------ slide 25
s = slide('Kết quả đã đạt được')
stat(s, [('4', 'vùng nguyên liệu có toạ độ và vùng trồng đã khoanh'),
         ('24', 'công đoạn số hoá theo đúng sơ đồ quy trình'),
         ('SHA-256', 'mã băm cho mỗi phiên bản công bố'),
         ('VBSN Besu', 'giao dịch neo đã xác nhận trên mạng thật')],
     top=TOP, height=Inches(1.55))
bullets(s, [
    ('Hệ thống chạy trên môi trường thật, không phải bản mô phỏng',
     'Toàn bộ luồng công bố, neo chuỗi và xác nhận giao dịch đã được kiểm '
     'chứng đầu cuối trên VBSN Besu.'),
    ('Kiểm chứng thực hiện được mà không cần tin vào hệ thống',
     'Script chain-verify đọc sự kiện trực tiếp từ blockchain và so với mã băm '
     'tự tính lại từ nội dung đang phục vụ.'),
    ('Chi phí vận hành gần như bằng không',
     'Không có máy chủ thường trực; mỗi lần neo chuỗi tốn 29.273 gas, tương '
     'đương 0,0000000293 VNX.'),
], top=Inches(3.95), size=14.5, gap=Inches(0.5))

# ------------------------------------------------------------------ slide 26
s = slide('Hạn chế và hướng phát triển')
label(s, ML, TOP - Inches(0.05), CW,
      'Hạn chế kỹ thuật · thuộc phạm vi phát triển tiếp')
cards(s, [
    ('Tài khoản và nhật ký',
     'Đang dùng một token quản trị chung, chưa có tài khoản riêng cho từng '
     'người vận hành và nhật ký từng lần sửa.'),
    ('Khối lượng và định mức',
     'Chưa ghi khối lượng vào và ra mỗi công đoạn, nên chưa kiểm được cân bằng '
     'khối lượng của mẻ theo ISO 22005.'),
    ('Tham số công đoạn',
     'Nhiệt độ, thời gian, tỷ lệ mới ở dạng mô tả tự do, chưa phải trường dữ '
     'liệu có cấu trúc.'),
], top=TOP + Inches(0.32), height=Inches(1.62))
label(s, ML, TOP + Inches(2.24), CW,
      'Điều kiện bên ngoài · phụ thuộc thủ tục của doanh nghiệp')
cards(s, [
    ('Mã GS1 (GTIN, GLN)',
     'Phụ thuộc thủ tục đăng ký mã doanh nghiệp với GS1 Việt Nam, không phải '
     'hạng mục lập trình. Phần hệ thống đã có phương án sẵn, xem trang 15.'),
    ('Chữ ký của từng nhà cung cấp',
     'Mỗi nhà cung cấp tự ký lô nguyên liệu bằng khoá riêng, hệ thống trả phí '
     'giao dịch. Cần từng đơn vị đồng ý giữ khoá; kế hoạch ở BLOCKCHAIN.md.'),
], top=TOP + Inches(2.61), height=Inches(1.62), cols=2)

# ------------------------------------------------------------------ slide 27
_n[0] += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=INK)
box(s, 0, 0, Inches(0.16), H, fill=ORANGE)
prod = os.path.join(ASSET, 'tule-products.jpg')
if os.path.exists(prod):
    s.shapes.add_picture(prod, Inches(7.85), Inches(1.9), Inches(4.6))
text(s, ML, Inches(2.15), Inches(6.5), Inches(1.4),
     'Truy xuất nguồn gốc chỉ có giá trị\nkhi người mua tự kiểm tra được',
     size=29, color=WHITE, bold=True, space=0, line=1.2)
box(s, ML, Inches(3.85), Inches(2.9), Pt(1.4), fill=RGBColor(0x3C, 0x5C, 0x4C))
text(s, ML, Inches(4.15), Inches(6.4), Inches(1.7),
     [('TÀI LIỆU KÈM THEO', {'size': 11, 'bold': True, 'color': ORANGE}),
      ('tai-lieu/KIEN_TRUC.md: kiến trúc, mô hình dữ liệu, giao diện lập '
       'trình', {'size': 13, 'color': PALE}),
      ('tai-lieu/HUONG_DAN_SU_DUNG.md: hướng dẫn vận hành',
       {'size': 13, 'color': PALE}),
      ('QUY_TRINH.md: đối chiếu quy trình sản xuất và tiêu chuẩn',
       {'size': 13, 'color': PALE})], space=5)
text(s, ML, Inches(6.15), Inches(6.4), Inches(0.4),
     'tule-trace.sontm.workers.dev', size=14, color=WHITE, bold=True, space=0)

prs.save(OUT)
print('Da tao:', OUT)
print('So slide:', len(prs.slides._sldIdLst))

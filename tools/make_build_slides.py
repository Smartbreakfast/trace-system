# -*- coding: utf-8 -*-
"""Bộ slide kỹ thuật: dựng hệ thống này như thế nào.

Khác `make_slides.py` (giới thiệu sản phẩm cho khách): bản này đi theo trục
kỹ thuật — phân tích nghiệp vụ, mô hình dữ liệu, rồi tách từng bộ phận của hai
sản phẩm (landing page và hệ truy xuất), mỗi bộ phận nói rõ dựng bằng gì, vì
sao chọn như vậy, và kết quả đo được.

    python tools/make_build_slides.py [thư mục ảnh]

Ảnh mặc định lấy ở `docs/ppt-ky-thuat/`.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'docs', 'ppt-ky-thuat')
OUT = os.path.join(ROOT, 'tai-lieu', 'Tu-Le-Trace-Ky-thuat-trien-khai.pptx')

INK = RGBColor(0x17, 0x35, 0x2A)
INK_SOFT = RGBColor(0x2C, 0x4C, 0x3E)
ORANGE = RGBColor(0xE2, 0x70, 0x3A)
ORANGE_INK = RGBColor(0xA8, 0x4B, 0x1B)
BG = RGBColor(0xF6, 0xF6, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xDA, 0xCC)
HAIR = RGBColor(0xE2, 0xE5, 0xDA)
MUTED = RGBColor(0x5C, 0x6B, 0x61)
GREEN_SOFT = RGBColor(0xEF, 0xF7, 0xEC)
GREEN_EDGE = RGBColor(0xCB, 0xE3, 0xC0)
TEAL = RGBColor(0x00, 0x7E, 0x88)

FONT = 'Segoe UI'
MONO = 'Consolas'
W, H = Inches(13.333), Inches(7.5)
ML = Inches(0.8)
CW = W - ML * 2

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------- tiện ích
def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    back = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    back.fill.solid()
    back.fill.fore_color.rgb = bg
    back.line.fill.background()
    back.shadow.inherit = False
    return s


def text(s, left, top, width, height, runs, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (content, opts) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opts.get('space', 6))
        p.line_spacing = opts.get('line', 1.2)
        run = p.add_run()
        run.text = content
        f = run.font
        f.name = opts.get('font', FONT)
        f.size = Pt(opts.get('size', 17))
        f.bold = opts.get('bold', False)
        f.color.rgb = opts.get('color', INK)
    return tb


def box(s, left, top, width, height, fill=WHITE, edge=LINE, radius=True, line_w=1.0):
    shape = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if edge is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = edge
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = 0.07
    return shape


def head(s, eyebrow, title, sub=None, width=None):
    """Tiêu đề slide. Tự chừa chỗ khi tiêu đề hoặc mô tả phải xuống dòng."""
    width = width or CW
    y = Inches(0.55)
    if eyebrow:
        text(s, ML, y, width, Inches(0.3),
             [(eyebrow.upper(), dict(size=12.5, bold=True, color=ORANGE_INK, space=0))])
        y += Inches(0.34)
    text(s, ML, y, width, Inches(0.9),
         [(title, dict(size=31, bold=True, color=INK, space=2))])
    per_line = max(int(width / Inches(0.1) * 0.40), 20)
    y += Inches(0.63) * max(1, -(-len(title) // per_line))
    if sub:
        lines = max(1, -(-len(sub) // max(int(width / Inches(0.1) * 0.76), 30)))
        text(s, ML, y, width, Inches(0.42) * lines,
             [(sub, dict(size=16, color=MUTED, line=1.32))])
        y += Inches(0.32) * lines + Inches(0.1)
    return y + Inches(0.14)


def shot(s, name, top, max_w=None, max_h=None, left=None, frame=True):
    from PIL import Image
    path = os.path.join(IMG, name)
    if not os.path.exists(path):
        print('  THIẾU ẢNH:', name)
        return None
    w, h = Image.open(path).size
    max_w = max_w or CW
    max_h = max_h or (H - top - Inches(0.5))
    scale = min(max_w / w, max_h / h)
    tw, th = int(w * scale), int(h * scale)
    x = left if left is not None else int((W - tw) / 2)
    if frame:
        f = box(s, Emu(x - 11000), Emu(top - 11000), Emu(tw + 22000), Emu(th + 22000),
                fill=WHITE, edge=LINE)
        f.shadow.inherit = False
    s.shapes.add_picture(path, Emu(x), Emu(top), width=Emu(tw), height=Emu(th))
    return Emu(top + th)


def table(s, rows, left, top, width, col_w, head_size=13, body_size=14.5, row_h=None):
    """Bảng phẳng: hàng đầu là tiêu đề, mỗi hàng một đường kẻ mảnh."""
    y = top
    for r, row in enumerate(rows):
        x = left
        height = row_h or (Inches(0.42) if r == 0 else Inches(0.52))
        for c, cell in enumerate(row):
            text(s, x, y + Inches(0.06), col_w[c] - Inches(0.2), height,
                 [(cell, dict(size=head_size if r == 0 else body_size,
                              bold=(r == 0),
                              color=MUTED if r == 0 else INK_SOFT,
                              space=0, line=1.22))])
            x += col_w[c]
        y += height
        if r == 0 or r < len(rows) - 1:
            rule = box(s, left, y, width, Pt(0.8), fill=HAIR if r else LINE,
                       edge=None, radius=False)
            rule.shadow.inherit = False
            y += Inches(0.08)
    return y


def card(s, left, top, width, height, title, lines, tint=WHITE, edge=LINE,
         title_color=INK, size=14.5):
    box(s, left, top, width, height, fill=tint, edge=edge)
    text(s, left + Inches(0.26), top + Inches(0.2), width - Inches(0.52), Inches(0.36),
         [(title, dict(size=16, bold=True, color=title_color, space=4))])
    runs = [(line, dict(size=size, color=MUTED, line=1.3, space=5)) for line in lines]
    text(s, left + Inches(0.26), top + Inches(0.62), width - Inches(0.52),
         height - Inches(0.85), runs)


def footer(s, note):
    text(s, ML, H - Inches(0.55), CW, Inches(0.3),
         [(note, dict(size=11.5, color=MUTED, space=0))])


def bullets(s, items, left, top, width, size=15.5, dot=ORANGE, gap=Inches(0.46)):
    y = top
    for item in items:
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, left, y + Inches(0.085), Pt(6), Pt(6))
        d.fill.solid()
        d.fill.fore_color.rgb = dot
        d.line.fill.background()
        d.shadow.inherit = False
        text(s, left + Inches(0.22), y, width - Inches(0.22), Inches(0.4),
             [(item, dict(size=size, color=INK_SOFT, line=1.28))])
        lines = max(1, int(len(item) / (width / Inches(0.098))) + 1)
        y += max(gap, Inches(0.27) * lines)
    return y


def section(title, index, note):
    """Slide phân đoạn, tách bốn phần của bài."""
    s = slide(INK)
    text(s, ML, Inches(2.6), CW, Inches(0.4),
         [('PHẦN %s' % index, dict(size=13, bold=True,
                                   color=RGBColor(0x8B, 0xBF, 0x76), space=6))])
    text(s, ML, Inches(3.05), Inches(10.5), Inches(1.0),
         [(title, dict(size=40, bold=True, color=WHITE, space=10))])
    text(s, ML, Inches(4.15), Inches(9.6), Inches(0.8),
         [(note, dict(size=17, color=RGBColor(0xA8, 0xC2, 0xB0), line=1.35))])


# =========================================================== A. mở đầu
def s_cover():
    s = slide(INK)
    text(s, ML, Inches(1.95), CW, Inches(0.4),
         [('BÁO CÁO KỸ THUẬT', dict(size=13.5, bold=True,
                                    color=RGBColor(0x8B, 0xBF, 0x76), space=0))])
    text(s, ML, Inches(2.35), Inches(11.6), Inches(2.0),
         [('Dựng hệ thống truy xuất',
           dict(size=44, bold=True, color=WHITE, line=1.14, space=0)),
          ('nguồn gốc như thế nào',
           dict(size=44, bold=True, color=WHITE, line=1.14, space=0))])
    text(s, ML, Inches(4.35), Inches(9.8), Inches(1.0),
         [('Phân tích nghiệp vụ, kiến trúc, và cách triển khai từng bộ phận của '
           'Tú Lệ Trace', dict(size=19, color=RGBColor(0xC4, 0xD6, 0xCA), line=1.35))])
    rule = box(s, ML, Inches(5.35), Inches(1.4), Pt(3), fill=ORANGE, edge=None, radius=False)
    rule.shadow.inherit = False
    text(s, ML, Inches(5.7), Inches(11.0), Inches(0.9),
         [('smartbreakfast.store  ·  trace.smartbreakfast.store  ·  Cloudflare Workers, D1, R2  ·  VBSN Besu 84001',
           dict(size=14.5, color=RGBColor(0x9F, 0xB8, 0xA8)))])


def s_pham_vi():
    s = slide()
    y = head(s, 'Phạm vi', 'Hai sản phẩm, hai bài toán khác nhau',
             'Cùng một thương hiệu nhưng hai nhóm yêu cầu kỹ thuật khác nhau, nên triển khai '
             'thành hai hệ thống riêng.')
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y, w, Inches(3.5), 'Landing page — smartbreakfast.store', [
        'Mục tiêu: giới thiệu sản phẩm, dẫn khách tới nơi mua.',
        'Nội dung tĩnh, đổi vài tuần một lần.',
        'Yêu cầu: SEO, chia sẻ mạng xã hội, tải nhanh.',
        'Không có dữ liệu động, không cần database.',
        '→ HTML/CSS/JS thuần trên Workers Static Assets.',
    ], size=14.5)
    card(s, ML + w + Inches(0.4), y, w, Inches(3.5),
         'Hệ truy xuất — trace.smartbreakfast.store', [
        'Mục tiêu: tra cứu hồ sơ một lô và kiểm chứng dữ liệu.',
        'Nội dung sinh theo từng lô, đổi mỗi lần công bố.',
        'Yêu cầu: toàn vẹn dữ liệu, phiên bản, bằng chứng trên chuỗi.',
        'Có hai nhóm người dùng với quyền khác nhau.',
        '→ Worker + D1 + R2, dựng HTML tại máy chủ, neo hash lên Besu.',
    ], tint=GREEN_SOFT, edge=GREEN_EDGE, title_color=RGBColor(0x2C, 0x6B, 0x3F), size=14.5)
    footer(s, 'Hai repo riêng: Smartbreakfast/trace-landing và Smartbreakfast/trace-system.')


def s_nghiep_vu():
    s = slide()
    y = head(s, 'Phân tích nghiệp vụ', 'Tác nhân, tình huống sử dụng, yêu cầu phi chức năng')
    rows = [
        ['Tác nhân', 'Tình huống sử dụng chính', 'Ràng buộc rút ra'],
        ['Người mua', 'Quét QR trên bao bì → đọc hồ sơ lô → tự kiểm chứng dữ liệu',
         'Mở trong vài giây trên 3G/4G yếu; không đăng nhập; đọc được trên điện thoại'],
        ['Người vận hành', 'Tạo lô, khai nguyên liệu và công đoạn, tải ảnh, công bố, in QR',
         'Một màn làm việc duy nhất cho mỗi lô; thao tác lặp lại hằng ngày phải ngắn'],
        ['Bên thứ ba\n(hội đồng, đối tác)', 'Kiểm chứng độc lập hồ sơ đã công bố',
         'Phải kiểm được mà không cần tin vào trang web của dự án'],
    ]
    y = table(s, rows, ML, y, CW, [Inches(2.4), Inches(4.6), CW - Inches(7.0)],
              row_h=Inches(0.66))
    y += Inches(0.18)
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y, w, Inches(1.65), 'Yêu cầu phi chức năng', [
        'Toàn vẹn: phát hiện mọi thay đổi sau khi công bố.',
        'Chi phí gần bằng 0 khi không có lượt quét; không máy chủ phải trực.',
    ], size=14)
    card(s, ML + w + Inches(0.4), y, w, Inches(1.65), 'Ngoài phạm vi bản này', [
        'Thanh toán và giỏ hàng — thuộc landing page.',
        'Tài khoản riêng cho từng nhân sự, chữ ký riêng của nhà cung cấp — ở lộ trình.',
    ], tint=BG, edge=LINE, size=14)


def s_quy_trinh():
    s = slide()
    y = head(s, 'Phân tích nghiệp vụ', 'Từ quy trình sản xuất thật tới mô hình dữ liệu',
             'Nhóm đọc thuyết minh quy trình của nhà sản xuất, tách thành bốn nhánh nguyên '
             'liệu và các công đoạn, rồi đối chiếu với mô hình dữ liệu đang có.')
    rows = [
        ['Trong quy trình thật', 'Phải có trong dữ liệu', 'Vì sao'],
        ['Bốn nguyên liệu đi bốn nhánh riêng rồi mới phối trộn',
         'ingredient_batches, mỗi nhánh một bản ghi',
         'Truy ngược từ thành phẩm về từng vùng trồng'],
        ['Mỗi công đoạn có khối lượng vào và ra',
         'input_quantity, output_quantity, quantity_unit',
         'Cho phép kiểm tra cân bằng khối lượng giữa các công đoạn'],
        ['Công đoạn có tham số kỹ thuật (nhiệt độ sấy, thời gian ngâm)',
         'params dạng JSON',
         'Phục vụ đối chiếu điều kiện chế biến khi kiểm tra'],
        ['Người trực tiếp làm từng công đoạn',
         'operator',
         'Ghi nhận trách nhiệm ở mức cá nhân'],
    ]
    table(s, rows, ML, y, CW, [Inches(4.5), Inches(4.0), CW - Inches(8.5)],
          row_h=Inches(0.66))
    footer(s, 'Hồ sơ số hoá hiện tại: 4 vùng nguyên liệu, 24 công đoạn nhánh + 2 công đoạn xưởng, 36 ảnh minh chứng mỗi lô.')


def s_mo_hinh():
    s = slide()
    y = head(s, 'Mô hình dữ liệu', 'Sáu bảng, ba quy tắc bất di bất dịch')
    shot(s, 'dg-du-lieu.png', y, max_h=Inches(4.6))
    footer(s, 'worker/schema.sql và worker/migrations — mọi thay đổi lược đồ đều đi qua một file migration có đánh số.')


# =========================================================== B. landing
def s_landing_kt():
    s = slide()
    y = head(s, 'Landing page', 'Yêu cầu và lựa chọn kỹ thuật',
             'Trang giới thiệu không có dữ liệu động; mỗi framework thêm vào đều làm tăng '
             'dung lượng tải lần đầu và thêm một bước build.')
    rows = [
        ['Hạng mục', 'Lựa chọn', 'Lý do'],
        ['Giao diện', 'HTML, CSS, JavaScript thuần', 'Không build step; sửa file là thấy ngay'],
        ['Phục vụ', 'Cloudflare Workers + Static Assets', 'Một lệnh deploy; miễn phí ở mức lưu lượng hiện tại'],
        ['Định tuyến', 'Worker đứng trước assets', 'Gộp http→https và www→apex vào một lần chuyển hướng'],
        ['Đường dẫn', 'Không đuôi .html, có 404 riêng', 'Địa chỉ sạch, chia sẻ dễ đọc'],
        ['SEO', 'sitemap.xml, robots.txt, thẻ Open Graph', 'Google đọc được, chia sẻ Zalo/Facebook có ảnh'],
    ]
    y2 = table(s, rows, ML, y, Inches(7.4), [Inches(1.7), Inches(2.5), Inches(3.2)],
               row_h=Inches(0.62))
    shot(s, 'l1-landing-home.png', y, max_w=Inches(4.7), max_h=Inches(3.6),
         left=int(ML + Inches(7.8)))
    footer(s, '43 file, 6 MB, không có node_modules trong sản phẩm cuối.')


def s_landing_code():
    s = slide()
    y = head(s, 'Landing page', 'Cấu hình phục vụ và hệ thống thiết kế')
    shot(s, 'code-landing-worker.png', y, max_w=Inches(6.0), max_h=Inches(2.9), left=int(ML))
    shot(s, 'code-landing-css.png', y, max_w=Inches(6.0), max_h=Inches(2.9),
         left=int(ML + Inches(6.3)))
    y2 = y + Inches(3.1)
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y2, w, Inches(1.5), 'src/index.js — 25 dòng', [
        'Worker chỉ làm hai việc chuyển hướng rồi trả phần còn lại cho assets.',
        'Cấu hình nằm trong repo thay vì trên dashboard, nên deploy bằng một lệnh và '
        'lịch sử thay đổi nằm trong git.',
    ])
    card(s, ML + w + Inches(0.4), y2, w, Inches(1.5), 'style.css — hệ thống thiết kế', [
        'Toàn bộ màu, cỡ chữ, khoảng cách khai báo thành biến CSS ở :root.',
        'Mỗi cặp màu đều ghi tỉ lệ tương phản đã kiểm theo chuẩn WCAG AA.',
    ])


def s_landing_ui():
    s = slide()
    y = head(s, 'Landing page', 'Thiết kế đáp ứng và tính nhất quán thương hiệu',
             'Bảng màu lấy từ chính bao bì hộp 600 g. Thanh điều hướng của hệ truy xuất '
             'dùng lại đúng bộ token này, để người quét QR không thấy mình rơi sang website khác.')
    shot(s, 'l2-landing-quytrinh.png', y, max_w=Inches(5.6), max_h=Inches(3.5), left=int(ML))
    shot(s, 'l4-landing-mobile.png', y, max_w=Inches(2.2), max_h=Inches(3.5),
         left=int(ML + Inches(5.9)))
    card(s, ML + Inches(8.4), y, CW - Inches(8.4), Inches(3.5), 'Điểm kỹ thuật', [
        'Breakpoint: 560px, 900px, 1220px.',
        'Thanh điều hướng thành nút Menu dưới 1220px.',
        'Ảnh nằm trong assets/img, không dùng CDN ngoài.',
        'Cùng bộ biến màu với trang truy xuất.',
    ], size=14)


# =========================================================== C. hệ truy xuất
def s_kien_truc():
    s = slide()
    y = head(s, 'Hệ truy xuất', 'Sơ đồ thành phần')
    shot(s, 'dg-thanh-phan.png', y, max_h=Inches(4.7))
    footer(s, 'Một Worker phục vụ cả API lẫn HTML; Pages đứng trước để nhận tên miền tuỳ chỉnh.')


def s_bp_du_lieu():
    s = slide()
    y = head(s, 'Bộ phận 1 — Tầng dữ liệu', 'D1 (SQLite) với lược đồ có phiên bản',
             'Mọi thay đổi lược đồ đều là một file migration đánh số, chạy được trên cả máy '
             'local lẫn production, không sửa tay bảng đang chạy.')
    shot(s, 'code-schema.png', y, max_w=Inches(7.3), max_h=Inches(3.2), left=int(ML))
    card(s, ML + Inches(7.6), y, CW - Inches(7.6), Inches(3.2), 'Quyết định thiết kế', [
        'Ngày tháng lưu dạng chuỗi dd.MM.yyyy đúng như người vận hành nhập, không đổi '
        'sang timestamp để giữ nguyên dữ liệu gốc.',
        'status chỉ có DRAFT và PUBLISHED; mọi trạng thái khác suy ra từ '
        'published_snapshots.',
        '9 migration tính tới hiện tại, gần nhất là bỏ cột bán kính vùng trồng.',
    ], size=14)


def s_bp_api():
    s = slide()
    y = head(s, 'Bộ phận 2 — Tầng API', 'Hono trên Worker, tách công khai và quản trị')
    shot(s, 'code-api.png', y, max_w=Inches(7.3), max_h=Inches(3.0), left=int(ML))
    card(s, ML + Inches(7.6), y, CW - Inches(7.6), Inches(3.0), 'Phân nhóm', [
        '/api/public/* — không cần xác thực, chỉ đọc bản đã công bố.',
        '/api/admin/* — chắn bằng Bearer token, so sánh theo thời gian hằng số '
        'để không lộ thông tin qua độ trễ.',
        '/api/media/* — phục vụ tệp R2, hỗ trợ header Range cho video.',
    ], size=14)
    y2 = y + Inches(3.2)
    rows = [
        ['Nhóm', 'Số route', 'Kiểm tra đầu vào'],
        ['Công khai', '6', 'Chuẩn hoá mã lô, giới hạn độ dài'],
        ['Quản trị', '18', 'Ép kiểu, giới hạn, kiểm toạ độ, kiểm định dạng tệp'],
    ]
    table(s, rows, ML, y2, Inches(8.0), [Inches(2.0), Inches(1.6), Inches(4.4)],
          row_h=Inches(0.48))


def s_bp_toan_ven():
    s = slide()
    y = head(s, 'Bộ phận 3 — Tầng toàn vẹn', 'Chuẩn hoá JSON rồi băm SHA-256',
             'Yêu cầu: cùng một hồ sơ phải luôn cho ra cùng một mã băm, không phụ thuộc '
             'thứ tự đọc dữ liệu hay môi trường chạy.')
    shot(s, 'code-hash.png', y, max_w=Inches(7.3), max_h=Inches(3.1), left=int(ML))
    card(s, ML + Inches(7.6), y, CW - Inches(7.6), Inches(3.1), 'Ba điều kiện', [
        'Khoá sắp theo bảng chữ cái, chuỗi chuẩn hoá Unicode NFC.',
        'Trường rỗng bị loại khỏi payload, để bổ sung trường mới không làm đổi mã băm '
        'của các lô đã công bố.',
        'Payload lưu đầy đủ chứ không chỉ mã băm, để dựng lại được nội dung bản cũ.',
    ], size=14)


def s_luong():
    s = slide()
    y = head(s, 'Bộ phận 3 — Tầng toàn vẹn', 'Luồng công bố và luồng kiểm chứng')
    shot(s, 'dg-luong.png', y, max_h=Inches(4.5))


def s_bp_chain():
    s = slide()
    y = head(s, 'Bộ phận 4 — Tích hợp blockchain', 'Ký và gửi bằng viem, có hàng chờ và cron')
    shot(s, 'code-chain.png', y, max_w=Inches(6.9), max_h=Inches(2.9), left=int(ML))
    shot(s, 'c1-giao-dich.png', y, max_w=Inches(5.0), max_h=Inches(2.9),
         left=int(ML + Inches(7.2)))
    y2 = y + Inches(3.1)
    rows = [
        ['Hạng mục', 'Giá trị thực tế'],
        ['Mạng', 'VBSN Besu, chainId 84001, RPC besu-rpc.vbsn.vn'],
        ['Hợp đồng', 'TuleTrace tại 0xcd6811f9…d3ca, triển khai ở khối 4.410.749'],
        ['Phí một lần neo', '0,0000057 VNX — xác nhận trong dưới 2 giây'],
        ['Cơ chế chịu lỗi', 'blockchain_outbox + cron 10 phút quét lại bản chưa gửi được'],
    ]
    table(s, rows, ML, y2, Inches(11.0), [Inches(3.0), Inches(8.0)], row_h=Inches(0.44))


def s_chain_bay():
    s = slide()
    y = head(s, 'Bộ phận 4 — Tích hợp blockchain', 'Hai điều kiện của mạng, gặp lúc triển khai')
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y, w, Inches(3.2), 'Máy ảo trước Shanghai', [
        'Triệu chứng: eth_estimateGas trả về "Invalid opcode: 0x5f".',
        'Nguyên nhân: solc 0.8.24 sinh opcode PUSH0, node chưa hỗ trợ.',
        'Xử lý: biên dịch với evmVersion "paris" trong scripts/deploy-contract.mjs.',
    ])
    card(s, ML + w + Inches(0.4), y, w, Inches(3.2), 'Giá gas tối thiểu của mạng', [
        'Triệu chứng: mọi giao dịch bị từ chối, thông báo chỉ có "RPC Request failed".',
        'Nguyên nhân: baseFeePerGas của mạng là 7 wei nên viem đặt trần 8 wei, '
        'trong khi node chỉ nhận từ 0,1 gwei.',
        'Xử lý: đọc eth_gasPrice rồi cộng biên 25% trước khi ký.',
    ], tint=GREEN_SOFT, edge=GREEN_EDGE, title_color=RGBColor(0x2C, 0x6B, 0x3F))
    text(s, ML, y + Inches(3.45), CW, Inches(0.7),
         [('Thông báo lỗi từ RPC không nêu nguyên nhân. Nhóm xác định bằng một script gọi '
           'trực tiếp RPC để đọc baseFeePerGas, gasPrice và nonce của ví.',
           dict(size=16, color=INK_SOFT, line=1.35))])


def s_hop_dong():
    s = slide()
    y = head(s, 'Bộ phận 4 — Tích hợp blockchain', 'Hợp đồng TuleTrace: chỉ phát sự kiện',
             'Hợp đồng không lưu nội dung hồ sơ. Nó khai báo ba sự kiện và một sổ vai trò, '
             'phần còn lại nằm ở D1 và R2.')
    shot(s, 'v2-code-contract.png', y, max_w=Inches(7.4), max_h=Inches(3.3), left=int(ML))
    card(s, ML + Inches(7.7), y, CW - Inches(7.7), Inches(3.3), 'Ba sự kiện', [
        'BatchAnchored — mã băm của lô thành phẩm, mã lô, số phiên bản, mốc thời gian.',
        'IngredientAnchored — dành cho bước nhà cung cấp tự ký.',
        'ActorSet — sổ vai trò: SUPPLIER, PACKAGER, BRAND.',
    ], size=14)
    footer(s, 'contracts/TuleTrace.sol · Solidity 0.8.24, biên dịch evmVersion "paris", triển khai một lần bằng scripts/deploy-contract.mjs.')


def s_bp_media():
    s = slide()
    y = head(s, 'Bộ phận 5 — Tệp minh chứng', 'R2 khoá theo nội dung, nén ở trình duyệt')
    shot(s, 'code-anh.png', y, max_w=Inches(7.0), max_h=Inches(2.5), left=int(ML))
    card(s, ML + Inches(7.3), y, CW - Inches(7.3), Inches(2.5), 'Vì sao nén lúc tải lên', [
        'Khoá R2 là sha256 của chính tệp, và khoá đó nằm trong payload đem đi băm.',
        'Nén lại sau khi công bố là làm sai lệch bản đã neo lên chuỗi.',
        'Nén trước khi gửi thì tệp được băm chính là tệp được phục vụ.',
    ], size=14)
    y2 = y + Inches(2.7)
    rows = [
        ['role', 'Hiển thị ở', 'Mức nén'],
        ['cover', 'Ảnh lớn mục Thông tin chung', '1400px, chất lượng 0,82'],
        ['gallery', 'Dải ảnh sản phẩm, ảnh công đoạn', '1280px, chất lượng 0,80'],
        ['area_map', 'Đầu hồ sơ nguyên liệu', '1600px, chất lượng 0,88'],
        ['certificate / lab_report', 'Mục Kiểm định chất lượng', '2000px, chất lượng 0,86'],
    ]
    table(s, rows, ML, y2, Inches(11.2), [Inches(3.2), Inches(4.6), Inches(3.4)],
          row_h=Inches(0.38))
    footer(s, 'Đo trên bộ ảnh thật của một lô: 9,07 MB xuống 5,19 MB. Ảnh nào nén lại không nhẹ hơn 12% thì giữ nguyên bản gốc.')


def s_bp_ssr():
    s = slide()
    y = head(s, 'Bộ phận 6 — Trang truy xuất', 'Trang truy xuất dựng HTML tại máy chủ',
             'Toàn bộ nội dung nằm trong HTML trả về từ Worker. Tắt JavaScript thì vẫn đổi '
             'mục được bằng liên kết, vẫn mở được chi tiết công đoạn bằng :target.')
    shot(s, 'p3-so-do.png', y, max_w=Inches(6.6), max_h=Inches(2.9), left=int(ML))
    shot(s, 'code-tree.png', y, max_w=Inches(5.2), max_h=Inches(2.9),
         left=int(ML + Inches(6.9)))
    y2 = y + Inches(3.05)
    text(s, ML, y2, CW, Inches(0.8),
         [('Sơ đồ quy trình là SVG sinh tại máy chủ từ danh sách công đoạn: không dùng thư viện '
           'vẽ đồ thị, không canvas, và bấm chọn được cả khi tắt JavaScript.',
           dict(size=16, color=INK_SOFT, line=1.35))])
    footer(s, 'worker/src/page/ — shell.ts, theme.ts, tree.ts, trace-page.ts, client.ts.')


def s_toi_uu_tai():
    s = slide()
    y = head(s, 'Bộ phận 6 — Trang truy xuất', 'Tối ưu lượt tải đầu tiên',
             'Ảnh của lô được yêu cầu ngay trong thẻ script ở đầu trang, không đợi phần còn '
             'lại của trang chạy xong.')
    shot(s, 'v1-code-index.png', y, max_w=Inches(7.6), max_h=Inches(3.3), left=int(ML))
    card(s, ML + Inches(7.9), y, CW - Inches(7.9), Inches(3.3), 'Kết quả đo', [
        'Ảnh bắt đầu tải ở 344 ms thay vì 1793 ms.',
        'Lượt tải sau đó của trang dùng lại bản trong cache, không tải lần hai.',
        'Font cắt gọn còn Latin + tiếng Việt, chuyển sang woff2, mỗi file 13-20 KB.',
        'CSS nhúng thẳng trong HTML, không thêm vòng đi về.',
    ], size=14)
    footer(s, 'web/index.html — đoạn script chạy trước, đọc hồ sơ lô rồi chèn thẻ preload cho ảnh bìa và dải ảnh.')


def s_bp_admin():
    s = slide()
    y = head(s, 'Bộ phận 7 — Màn quản trị', 'Flutter Web, một màn làm việc cho mỗi lô')
    shot(s, 'a3-man-lam-viec.png', y, max_w=Inches(6.2), max_h=Inches(3.3), left=int(ML))
    shot(s, 'a9-lich-su-qr.png', y, max_w=Inches(6.2), max_h=Inches(3.3),
         left=int(ML + Inches(6.5)))
    y2 = y + Inches(3.5)
    text(s, ML, y2, CW, Inches(0.9),
         [('Bốn khối theo đúng thứ tự công việc: thông tin lô → ảnh và hồ sơ → nguyên liệu và '
           'công đoạn → hoàn thành lô. Khối cuối hiện trạng thái toàn vẹn, mã băm, lịch sử các '
           'bản đã công bố kèm tình trạng giao dịch, và mã QR tải về dạng PNG cho bản in bao bì.',
           dict(size=16, color=INK_SOFT, line=1.35))])
    footer(s, 'Màn quản trị giữ Flutter: tải một lần rồi nằm trong cache trình duyệt, và dùng lại được khi phát triển ứng dụng di động.')


def s_bp_ha_tang():
    s = slide()
    y = head(s, 'Bộ phận 8 — Hạ tầng và triển khai', 'Một lệnh build, hai lệnh deploy',
             'Không có máy chủ phải cài đặt hay trực. Cấu hình nằm trong repo dưới dạng file.')
    rows = [
        ['Thành phần', 'Cấu hình', 'Ghi chú'],
        ['Trang truy xuất + API', 'wrangler deploy', 'Worker mang theo D1, R2, biến môi trường, cron'],
        ['Bản build Flutter', 'flutter build web --pwa-strategy=none', 'Phục vụ ở /admin'],
        ['Tên miền', 'CNAME trace → Pages, DNS only', 'Pages chuyển tiếp /, /t/* và /api/* về Worker'],
        ['Tác vụ định kỳ', 'cron "*/10 * * * *"', 'Quét lại hàng chờ neo chuỗi khi mạng chập'],
    ]
    y2 = table(s, rows, ML, y, Inches(11.5), [Inches(3.0), Inches(3.8), Inches(4.7)],
               row_h=Inches(0.46))
    y2 += Inches(0.24)
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y2, w, Inches(2.4), 'Quản lý bí mật', [
        'PRIVATE_KEY và ADMIN_TOKEN đặt bằng wrangler secret put.',
        'Máy phát triển dùng worker/.dev.vars, đã đưa vào .gitignore.',
        'Repo chỉ có .dev.vars.example với giá trị mẫu.',
    ], size=13.5)
    card(s, ML + w + Inches(0.4), y2, w, Inches(2.4), 'Chi phí vận hành', [
        'Workers, D1, R2 nằm trong mức miễn phí ở lưu lượng hiện tại.',
        'R2 không tính phí truyền dữ liệu ra.',
        'Phí chuỗi: 0,0000057 VNX cho mỗi lần công bố.',
    ], tint=BG, size=13.5)


# =========================================================== D. chất lượng
def s_hieu_nang():
    s = slide()
    y = head(s, 'Đo lường', 'Hiệu năng trang truy xuất: trước và sau khi bỏ Flutter',
             'Bản đầu dùng Flutter Web cho cả trang công khai. Đo lại trên cùng một lô, '
             'cùng một máy, mô phỏng điện thoại phổ thông với mạng 4G yếu.')
    shot(s, 'fig-toc-do.png', y, max_h=Inches(3.5))
    footer(s, 'Phương pháp đo: Chrome headless, CPU throttling 4x, mạng 1,6 Mbps / độ trễ 150 ms, cache trống.')


def s_kiem_thu():
    s = slide()
    y = head(s, 'Chất lượng', 'Kiểm thử và kiểm chứng độc lập')
    w = (CW - Inches(0.5)) / 3
    card(s, ML, y, w, Inches(3.0), '38 widget test (Flutter)', [
        'Bao gồm các ranh giới quan trọng: không rơi về dữ liệu mẫu khi API hỏng, '
        'chỉ hiện chữ "blockchain" khi thật sự có giao dịch, sơ đồ dựng đúng thứ tự.',
    ])
    card(s, ML + w + Inches(0.25), y, w, Inches(3.0), 'Smoke test API', [
        'Chạy thẳng vào Worker: tạo lô, nhập liệu, công bố, kiểm tra mã băm đổi đúng lúc '
        'và giữ nguyên khi thêm trường rỗng.',
    ])
    card(s, ML + (w + Inches(0.25)) * 2, y, w, Inches(3.0), 'Kiểm chứng từ chuỗi', [
        'scripts/chain-verify.mjs đọc sự kiện BatchAnchored thẳng từ node Besu, '
        'không qua API của hệ thống.',
        'Phép kiểm chứng này không phụ thuộc vào API của hệ thống.',
    ], tint=GREEN_SOFT, edge=GREEN_EDGE, title_color=RGBColor(0x2C, 0x6B, 0x3F))
    text(s, ML, y + Inches(3.25), CW, Inches(0.8),
         [('node scripts/chain-verify.mjs 0xcd6811f9a06d706978033edb6ecaf72d37ffd3ca <sha256>',
           dict(size=15, font=MONO, color=INK, space=6)),
          ('Kết quả in ra mã lô, số phiên bản, ví đã ký, thời điểm và số khối.',
           dict(size=15, color=MUTED))])


def s_cong_cu():
    s = slide()
    y = head(s, 'Quy trình phát triển', 'Công cụ sử dụng và cách phân chia công việc')
    rows = [
        ['Công cụ', 'Dùng vào việc'],
        ['Visual Studio Code', 'Viết mã, chạy thử, xem log'],
        ['Git và GitHub', 'Lưu lịch sử thay đổi, hai repo riêng cho landing và hệ truy xuất'],
        ['Wrangler CLI', 'Chạy Worker ở máy local, deploy, đặt biến bí mật, chạy migration D1'],
        ['Trợ lý AI (Claude)', 'Hỗ trợ viết mã, giải thích lỗi, rà soát lại đoạn đã viết'],
        ['Chrome DevTools', 'Đo thời gian tải, kiểm tra request và kích thước tệp'],
    ]
    y2 = table(s, rows, ML, y, Inches(11.5), [Inches(3.4), Inches(8.1)], row_h=Inches(0.44))
    y2 += Inches(0.25)
    w = (CW - Inches(0.4)) / 2
    card(s, ML, y2, w, Inches(2.2), 'AI hỗ trợ phần nào', [
        'Sinh nhanh đoạn mã lặp lại, gợi ý cú pháp thư viện chưa quen.',
        'Giải thích thông báo lỗi và đề xuất hướng xử lý.',
        'Rà soát mã đã viết, chỉ ra trường hợp biên chưa xử lý.',
    ], size=14)
    card(s, ML + w + Inches(0.4), y2, w, Inches(2.2), 'Nhóm tự quyết định', [
        'Phạm vi dữ liệu công bố và dữ liệu giữ nội bộ.',
        'Mô hình dữ liệu và ràng buộc thiết kế RB-01…RB-07.',
        'Đối chiếu quy trình sản xuất thật với từng công đoạn.',
        'Thử nghiệm trên thiết bị thật và xử lý phản hồi.',
    ], tint=GREEN_SOFT, edge=GREEN_EDGE, title_color=RGBColor(0x2C, 0x6B, 0x3F), size=14)


def s_rang_buoc():
    s = slide()
    y = head(s, 'Chất lượng', 'Ràng buộc thiết kế được kiểm soát trong mã nguồn')
    rows = [
        ['Mã', 'Ràng buộc'],
        ['RB-01', 'Không bao giờ lặng lẽ rơi về dữ liệu mẫu khi API hỏng'],
        ['RB-02', 'Chỉ hiển thị trạng thái blockchain khi thật sự có giao dịch'],
        ['RB-03', 'Trường không có dữ liệu không được đưa vào payload công bố'],
        ['RB-04', 'Không xoá object R2 khi còn bản công bố tham chiếu tới khoá đó'],
        ['RB-05', 'Không hiển thị vùng nguyên liệu chính xác hơn mức dữ liệu thật sự có'],
        ['RB-06', 'Dữ liệu chi phí, giá vốn, lợi nhuận không được đưa vào payload công bố'],
        ['RB-07', 'Công bố không ghi đè bản trước; mỗi lần công bố tạo một phiên bản mới'],
    ]
    table(s, rows, ML, y, Inches(11.4), [Inches(1.3), Inches(10.1)], row_h=Inches(0.46))
    footer(s, 'Danh sách đầy đủ kèm lý do: tai-lieu/KIEN_TRUC.md, mục Ràng buộc thiết kế.')


def s_ket_qua():
    s = slide(INK)
    text(s, ML, Inches(0.95), CW, Inches(0.4),
         [('KẾT QUẢ VÀ LỘ TRÌNH', dict(size=13, bold=True,
                                       color=RGBColor(0x8B, 0xBF, 0x76), space=0))])
    text(s, ML, Inches(1.45), Inches(11.4), Inches(0.8),
         [('Đang chạy trên môi trường thật', dict(size=36, bold=True, color=WHITE, space=0))])

    stats = [('4', 'lô đã công bố'), ('26', 'công đoạn mỗi lô'),
             ('36', 'ảnh minh chứng mỗi lô'), ('1,0 s', 'nội dung hiện ra, 4G yếu')]
    w = (CW - Inches(0.75)) / 4
    for i, (num, label_text) in enumerate(stats):
        left = ML + (w + Inches(0.25)) * i
        b = box(s, left, Inches(2.5), w, Inches(1.35), fill=RGBColor(0x1E, 0x42, 0x34),
                edge=RGBColor(0x2C, 0x5A, 0x46))
        b.shadow.inherit = False
        text(s, left + Inches(0.25), Inches(2.68), w - Inches(0.5), Inches(0.5),
             [(num, dict(size=30, bold=True, color=WHITE, space=2))])
        text(s, left + Inches(0.25), Inches(3.22), w - Inches(0.5), Inches(0.4),
             [(label_text, dict(size=13.5, color=RGBColor(0x9F, 0xB8, 0xA8)))])

    text(s, ML, Inches(4.2), Inches(11.4), Inches(0.4),
         [('Lộ trình', dict(size=20, bold=True, color=RGBColor(0xC4, 0xD6, 0xCA), space=8))])
    items = [
        'Tài khoản riêng cho từng người vận hành, kèm nhật ký thao tác — 2 tuần',
        'Ứng dụng di động cho người nhập liệu tại xưởng — 3 tuần',
        'Mỗi nhà cung cấp tự ký công đoạn của mình — đang thiết kế, xem BLOCKCHAIN.md',
        'Mở rộng lên 100 lô — 1 tháng',
        'Bổ sung mã GS1 (GTIN, GLN) để liên thông hệ thống truy xuất quốc gia — 2 tuần',
    ]
    y = Inches(4.75)
    for item in items:
        text(s, ML, y, Inches(11.4), Inches(0.35),
             [('—  ' + item, dict(size=15, color=RGBColor(0xA8, 0xC2, 0xB0), space=0))])
        y += Inches(0.42)


for build in [
    s_cover, s_pham_vi, s_nghiep_vu, s_quy_trinh, s_mo_hinh,
    lambda: section('Landing page', 'B', 'Trang giới thiệu: yêu cầu, lựa chọn kỹ thuật, hệ thống thiết kế'),
    s_landing_kt, s_landing_code, s_landing_ui,
    lambda: section('Hệ truy xuất', 'C', 'Tám bộ phận: dữ liệu, API, toàn vẹn, blockchain, tệp, trang truy xuất, quản trị, hạ tầng'),
    s_kien_truc, s_bp_du_lieu, s_bp_api, s_bp_toan_ven, s_luong,
    s_bp_chain, s_hop_dong, s_chain_bay, s_bp_media, s_bp_ssr, s_toi_uu_tai,
    s_bp_admin, s_bp_ha_tang,
    lambda: section('Đo lường và chất lượng', 'D', 'Hiệu năng, quy trình phát triển, kiểm thử, ràng buộc thiết kế'),
    s_hieu_nang, s_cong_cu, s_kiem_thu, s_rang_buoc, s_ket_qua,
]:
    build()

prs.save(OUT)
print('Đã lưu', OUT)
